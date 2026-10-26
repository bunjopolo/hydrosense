"""CSV and Excel imports; preserve tag names exactly."""
import base64
import io
import zipfile
from pathlib import Path

from .model import read_csv


def read_upload(data, config):
    timestamp = data.get('timestamp_column') if (data.get('purpose') == 'inspect' or data.get('auto_configure')) else config['timestamp_column']
    timestamp = timestamp or config['timestamp_column']
    source_timezone = data.get('source_timezone', config.get('source_timezone'))
    if 'xlsx' in data:
        import openpyxl
        raw = base64.b64decode(data['xlsx'], validate=True)
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 256 * 1024 * 1024:
                raise ValueError('Expanded workbook exceeds 256 MB. Export the relevant sheet as CSV.')
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            name = data.get('sheet') or workbook.sheetnames[0]
            if name not in workbook.sheetnames:
                raise ValueError('Excel sheet not found. Available: ' + ', '.join(workbook.sheetnames))
            rows = workbook[name].iter_rows(values_only=True)
            import csv
            stream = io.StringIO()
            writer = csv.writer(stream)
            for row in rows:
                writer.writerow(row)
            text = stream.getvalue()
        finally:
            workbook.close()
    else:
        text = data['csv']
    return read_csv(text, timestamp, source_timezone)


def read_file(path, config, sheet=''):
    path = Path(path)
    data = {'xlsx': base64.b64encode(path.read_bytes()).decode()} if path.suffix.lower() == '.xlsx' else {'csv': path.read_text(encoding='utf-8-sig')}
    return read_upload({**data, 'sheet': sheet}, config)
